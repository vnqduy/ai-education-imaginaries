#!/usr/bin/env python3
"""Reproduce current reviewed SIPC summaries without network or LLM calls."""
import argparse, csv, hashlib, json, tempfile
from collections import Counter
from pathlib import Path
import openpyxl
ROOT = Path(__file__).resolve().parents[1]
GROUPS = ('public/central', 'commercial/general')
def read(path):
    with path.open(encoding='utf-8-sig', newline='') as f:
        return list(csv.DictReader(f))
def share(rows, col):
    valid = [r for r in rows if r[col.split('__')[0]+'__status'] == 'coded']
    n = sum(r[col] == '1' for r in valid)
    return dict(n=n, N=len(valid), pct=100*n/len(valid) if valid else None)
def compare(rows, cols):
    result = {}
    for col in cols:
        a, b = [share([r for r in rows if r['outlet_type']==g], col) for g in GROUPS]
        gaps = []
        for domain in sorted({r['outlet_domain'] for r in rows}):
            av, bv = [share([r for r in rows if r['outlet_type']==g and r['outlet_domain']!=domain], col) for g in GROUPS]
            if av['pct'] is not None and bv['pct'] is not None:
                gaps.append(av['pct']-bv['pct'])
        result[col] = dict(A=a, B=b, gap=a['pct']-b['pct'] if a['pct'] is not None and b['pct'] is not None else None, loo_min=min(gaps) if gaps else None, loo_max=max(gaps) if gaps else None)
    return result

def main():
    ap=argparse.ArgumentParser(description=__doc__)
    ap.add_argument('--output-dir', type=Path, default=Path(tempfile.gettempdir())/'sipc-reviewed-results')
    out=ap.parse_args().output_dir; out.mkdir(parents=True, exist_ok=True)
    rows=read(ROOT/'data/final_analysis_dataset.csv')
    cols=[c for c in rows[0] if '__' in c and not c.endswith('__status')]
    assert len(rows)==917 and len({r['article_id'] for r in rows})==917 and len(cols)==89
    corrections=read(ROOT/'evidence/coding_corrections.csv')
    assert len(corrections)==31 and len({c['article_id'] for c in corrections})==26
    wb=openpyxl.load_workbook(ROOT/'data/final_article_coding.xlsx',read_only=True,data_only=True)
    it=iter(wb['Coded articles'].values); headers=next(it)
    frozen={r['article_id']:r for values in it if (r:=dict(zip(headers,values)))['id_match_status']=='matched'}
    wb.close()
    ledger={(r['article_id'],r['field']):r for r in corrections}
    seen=set()
    for r in rows:
        original=frozen[r['article_id']]
        for col in [c for c in r if '__' in c]:
            old='' if original[col] is None else str(original[col])
            if old != r[col]:
                key=(r['article_id'],col); change=ledger[key]
                assert old==change['old_value'] and r[col]==change['new_value']
                seen.add(key)
        for col in cols:
            assert r[col] in ('0','1') if r[col.split('__')[0]+'__status']=='coded' else r[col]==''
    assert seen==set(ledger)
    post=[r for r in rows if r['date']>='2022-11-30']
    assert len(post)==871 and Counter(r['outlet_type'] for r in post)==dict(zip(GROUPS,[433,438]))
    comparisons=compare(post,cols)
    (out/'outlet_comparison.json').write_text(json.dumps(comparisons,indent=2)+'\n')
    # One workbook consolidates descriptive tables and sensitivity results.
    book=openpyxl.Workbook(); book.remove(book.active)
    overall=book.create_sheet('Overall');overall.append(['indicator','n','valid_N','percent'])
    for col in cols:
        s=share(rows,col);overall.append([col,s['n'],s['N'],s['pct']])
    for name, selected in [('Outlet comparison',post),('Yes-only',[r for r in post if r['sipc_presence']=='yes']),('January cutoff',[r for r in rows if r['date']>='2023-01-01']),('Within generative AI',[r for r in post if r['technology_object__generative']=='1'])]:
        sheet=book.create_sheet(name);sheet.append(['indicator','A_n','A_valid_N','A_percent','B_n','B_valid_N','B_percent','gap_pp','leave_one_out_min','leave_one_out_max'])
        for col,v in compare(selected,cols).items():
            sheet.append([col,v['A']['n'],v['A']['N'],v['A']['pct'],v['B']['n'],v['B']['N'],v['B']['pct'],v['gap'],v['loo_min'],v['loo_max']])
    corpus=book.create_sheet('Corpus composition');corpus.append(['outlet_type','period','articles'])
    for (g,p),n in sorted(Counter((r['outlet_type'],r['chatgpt_period']) for r in rows).items()):corpus.append([g,p,n])
    for sheet in book:
        sheet.freeze_panes='A2';sheet.auto_filter.ref=sheet.dimensions
        sheet.column_dimensions['A'].width=43
        for cells in list(sheet.rows)[1:]:
            for cell in cells:
                if isinstance(cell.value,float):cell.number_format='0.00'
    book.save(out/'results.xlsx')
    files=['data/final_analysis_dataset.csv','data/final_article_coding.xlsx','evidence/coding_corrections.csv','methods/coding_codebook.xlsx']
    audit=dict(validation='passed',retained_articles=len(rows),post_articles=len(post),indicators=len(cols),reviewed_cells=len(seen),input_sha256={p:hashlib.sha256((ROOT/p).read_bytes()).hexdigest() for p in files},scope='Verifies approved corrections against frozen coding and reproduces numerical summaries; does not regenerate extraction or qualitative interpretation.')
    (out/'validation.json').write_text(json.dumps(audit,indent=2)+'\n')
    print(json.dumps({k:v for k,v in audit.items() if k!='input_sha256'},indent=2))
if __name__=='__main__':main()
